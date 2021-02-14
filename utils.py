import configparser
import csv
import datetime
import time

from docx import Document
from pynput.keyboard import Controller as KB_Controller
from pynput.keyboard import Events, Key
from pynput.mouse import Button
from pynput.mouse import Controller as Mouse_Controller

from mysql_utils import connect, insert_records


def get_period(time):
    """Gets period based on time"""
    periods = {
        1: (datetime.time(7, 30), datetime.time(8, 40)),
        2: (datetime.time(8, 40), datetime.time(9, 30)),
        3: (datetime.time(9, 30), datetime.time(10, 25)),
        4: (datetime.time(10, 25), datetime.time(11, 15)),
        5: (datetime.time(11, 15), datetime.time(12, 15)),
    }
    for period in periods:
        if periods[period][0] <= time.time() < periods[period][1]:
            return period
    print("No class at the moment")
    input("press Enter to exit...")
    raise SystemExit


def get_subject():
    """Retrieves subject based on timetable"""
    with open("timetable.csv", "r") as timetable:
        period = get_period(datetime.datetime.now())
        timetableReader = list(csv.reader(timetable))
    day = int(datetime.datetime.now().strftime("%w")) + 1
    subject = timetableReader[day][period]
    print("Day:", datetime.datetime.now().strftime("%A"))
    print(f"Period : {period}")
    print(f"subject : {subject}")
    return subject


def get_credentials(subject):
    """Retrieves id and password corresponding to subject"""
    subject = subject.upper()
    with open("passwords.csv", "r") as passwords:
        reader = csv.DictReader(passwords)
        for row in reader:
            if row["subject"].upper() == subject:
                return row["id"], row["password"]
        else:
            print("Failed to obtain id and password")
            raise SystemExit


def auto_type(id, password, joinposn):
    """Executes actions on Zoom.exe"""

    print("id:", id)
    keyboard = KB_Controller()
    mouse = Mouse_Controller()
    with keyboard.pressed(Key.alt_l):
        keyboard.press(Key.tab)
        keyboard.release(Key.tab)
    time.sleep(0.5)  #
    mouse.position = joinposn
    mouse.click(Button.left)
    time.sleep(1)  #
    keyboard.type(id)
    keyboard.press(Key.enter)
    keyboard.release(Key.enter)
    time.sleep(3.5)  #
    keyboard.type(password)
    keyboard.press(Key.enter)
    keyboard.release(Key.enter)


def append(values):
    conn, cursor = connect()
    insert_records(cursor, values)
    conn.commit()


def reset():
    conn, cursor = connect()
    cursor.execute("DROP DATABASE ZOOMER;")
    conn.commit()
    print("Database deleted.")


def change_pass(values):
    subject = values[0].upper()
    password = values[1]
    with open("passwords.csv", "r") as passwords:
        reader = list(csv.reader(passwords))
        for row in reader:
            if row[0].upper() == subject:
                row[2] = password
                break

    with open("passwords.csv", "w") as asd:
        for line in reader:
            asd.write(f"{line[0]},{line[1]},{line[2]}\n")
        raise SystemExit


def update_pass(path_to_new_pass):
    """updates passwords with passwords extracted from a word doc."""
    conn, cursor = connect()
    doc = Document(path_to_new_pass)
    for table in doc.tables:
        for row in range(1, 8):
            subject = table.cell(row, 2).text.strip()
            id = table.cell(row, 3).text.strip()
            password = table.cell(row, 4).text.strip()
            records = (subject, id, password)
            insert_records(cursor, records)
    conn.commit()
    raise SystemExit


def getjoinposn():
    mouse = Mouse_Controller()
    print(
        "Open zoom and hover the mouse cursor "
        'above the "join" button and press <CTRL>.'
    )

    def get_coord():
        with Events() as events:
            for event in events:
                if event.key == Key.ctrl_l:
                    posn = mouse.position
                    return str(posn)[1:-1]

    config = configparser.ConfigParser()
    config.read("data.ini")
    config.set("VALUES", "JOIN", get_coord())

    with open("data.ini", "w") as data:
        config.write(data)
    print("\n>Mouse position successfully recorded!")
    input(">Press enter to exit...")
    raise SystemExit
